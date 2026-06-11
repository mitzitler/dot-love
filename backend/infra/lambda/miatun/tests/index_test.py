import io
import json
import os
import sys
import unittest
from unittest.mock import MagicMock, patch

# Set required environment variables before importing index
os.environ["user_table_name"] = "test_user_table"
os.environ["registry_item_table_name"] = "test_registry_item_table"
os.environ["registry_claim_table_name"] = "test_registry_claim_table"
os.environ["export_bucket_name"] = "test_export_bucket"
os.environ["internal_api_key"] = "test_internal_key"
os.environ["AWS_DEFAULT_REGION"] = "us-east-1"

# Add parent directory to path to import the lambda function
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

# Mock boto3 before importing index to avoid AWS credentials requirement
with patch("boto3.client"):
    import index

from openpyxl import load_workbook


def make_item(item_id, item_name, claimant_id="", brand="", price_cents=None):
    item = {
        "id": item_id,
        "item_name": item_name,
        "brand": brand,
        "claimant_id": claimant_id,
        "received": True,
    }
    if price_cents is not None:
        item["price_cents"] = price_cents
    return item


def make_claim(item_id, claimant_id, claim_state="CLAIMED"):
    return {
        "id": f"claim-{item_id}-{claimant_id}",
        "item_id": item_id,
        "claimant_id": claimant_id,
        "claim_state": claim_state,
    }


def make_user(first_last, pair="", **contact):
    user = {
        "first_last": first_last,
        "guest_pair_first_last": pair,
        "phone": "+15555550100",
        "street": "123 Main St",
        "second_line": "Apt 4",
        "city": "Brooklyn",
        "state_loc": "NY",
        "zipcode": "11201",
        "country": "USA",
    }
    user.update(contact)
    return user


class TestHelpers(unittest.TestCase):
    """Test suite for formatting helpers."""

    def test_title_case_first_last(self):
        self.assertEqual(index.title_case_first_last("john_smith"), "John Smith")

    def test_title_case_single_token(self):
        self.assertEqual(index.title_case_first_last("plus_one"), "Plus One")
        self.assertEqual(index.title_case_first_last("cher"), "Cher")

    def test_title_case_empty(self):
        self.assertEqual(index.title_case_first_last(""), "")
        self.assertEqual(index.title_case_first_last(None), "")

    def test_format_price_cents(self):
        self.assertEqual(index.format_price_cents(12999), "$129.99")
        self.assertEqual(index.format_price_cents(0), "$0.00")
        self.assertEqual(index.format_price_cents(None), "")
        self.assertEqual(index.format_price_cents("garbage"), "")


class TestBuildThankYouRows(unittest.TestCase):
    """Test suite for the claims/items/users join logic."""

    def test_full_row_with_partner(self):
        items = {"i1": make_item("i1", "Stand Mixer", brand="KitchenAid", price_cents=44999)}
        claims = [make_claim("i1", "john_smith")]
        users = {"john_smith": make_user("john_smith", pair="jane_smith")}

        rows = index.build_thank_you_rows(items, claims, users)

        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row.guest, "John Smith")
        self.assertEqual(row.partner, "Jane Smith")
        self.assertEqual(row.gift, "Stand Mixer")
        self.assertEqual(row.brand, "KitchenAid")
        self.assertEqual(row.price, "$449.99")
        self.assertEqual(row.phone, "+15555550100")
        self.assertEqual(row.street, "123 Main St")
        self.assertEqual(row.city, "Brooklyn")

    def test_claim_for_non_received_item_excluded(self):
        items = {}  # nothing received
        claims = [make_claim("i1", "john_smith")]

        rows = index.build_thank_you_rows(items, claims, {})

        self.assertEqual(rows, [])

    def test_unclaimed_claim_state_excluded(self):
        items = {"i1": make_item("i1", "Vase")}
        claims = [make_claim("i1", "john_smith", claim_state="UNCLAIMED")]

        rows = index.build_thank_you_rows(items, claims, {})

        # The UNCLAIMED claim is skipped, but the received item still appears
        # via the fallback (using the item's own claimant_id, blank here).
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].guest, "")

    def test_purchased_claim_state_included(self):
        items = {"i1": make_item("i1", "Vase")}
        claims = [make_claim("i1", "john_smith", claim_state="PURCHASED")]

        rows = index.build_thank_you_rows(items, claims, {})

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].guest, "John Smith")

    def test_claimant_missing_from_users_table(self):
        items = {"i1": make_item("i1", "Blender")}
        claims = [make_claim("i1", "plus_one")]

        rows = index.build_thank_you_rows(items, claims, {})

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].guest, "Plus One")
        self.assertEqual(rows[0].partner, "")
        self.assertEqual(rows[0].phone, "")
        self.assertEqual(rows[0].street, "")

    def test_duplicate_claims_deduped(self):
        items = {"i1": make_item("i1", "Toaster")}
        claims = [
            make_claim("i1", "john_smith"),
            make_claim("i1", "john_smith", claim_state="PURCHASED"),
        ]

        rows = index.build_thank_you_rows(items, claims, {})

        self.assertEqual(len(rows), 1)

    def test_received_item_without_claim_falls_back_to_item_claimant(self):
        items = {"i1": make_item("i1", "Wok", claimant_id="amy_pond")}
        users = {"amy_pond": make_user("amy_pond")}

        rows = index.build_thank_you_rows(items, [], users)

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0].guest, "Amy Pond")
        self.assertEqual(rows[0].phone, "+15555550100")

    def test_rows_sorted_by_guest_then_gift(self):
        items = {
            "i1": make_item("i1", "Zester"),
            "i2": make_item("i2", "Apron"),
            "i3": make_item("i3", "Kettle"),
        }
        claims = [
            make_claim("i1", "zoe_washburne"),
            make_claim("i2", "amy_pond"),
            make_claim("i3", "amy_pond"),
        ]

        rows = index.build_thank_you_rows(items, claims, {})

        self.assertEqual(
            [(r.guest, r.gift) for r in rows],
            [("Amy Pond", "Apron"), ("Amy Pond", "Kettle"), ("Zoe Washburne", "Zester")],
        )


class TestBuildXlsx(unittest.TestCase):
    """Test suite for xlsx generation."""

    def test_workbook_round_trip(self):
        items = {"i1": make_item("i1", "Stand Mixer", brand="KitchenAid", price_cents=44999)}
        claims = [make_claim("i1", "john_smith")]
        users = {"john_smith": make_user("john_smith", pair="jane_smith")}
        rows = index.build_thank_you_rows(items, claims, users)

        xlsx_bytes = index.build_xlsx(rows)
        workbook = load_workbook(io.BytesIO(xlsx_bytes))
        sheet = workbook.active

        self.assertEqual(sheet.title, "Thank You Cards")
        self.assertEqual(sheet.freeze_panes, "A2")
        header = tuple(cell.value for cell in sheet[1])
        self.assertEqual(header, index.SHEET_HEADERS)
        self.assertTrue(sheet[1][0].font.bold)
        data_row = tuple(cell.value for cell in sheet[2])
        self.assertEqual(data_row[0], "John Smith")
        self.assertEqual(data_row[1], "Jane Smith")
        self.assertEqual(data_row[2], "Stand Mixer")

    def test_empty_rows_still_produces_header(self):
        xlsx_bytes = index.build_xlsx([])
        sheet = load_workbook(io.BytesIO(xlsx_bytes)).active
        self.assertEqual(sheet.max_row, 1)


class TestExportEndpoint(unittest.TestCase):
    """Test suite for the API routes through the full handler."""

    @staticmethod
    def make_event(path, method="GET", headers=None):
        return {
            "version": "2.0",
            "rawPath": path,
            "rawQueryString": "",
            "headers": headers or {},
            "requestContext": {
                "http": {"method": method, "path": path, "sourceIp": "127.0.0.1"},
                "requestId": "test-request-id",
                "stage": "$default",
            },
            "isBase64Encoded": False,
        }

    @staticmethod
    def make_context():
        ctx = MagicMock()
        ctx.function_name = "miatun"
        ctx.memory_limit_in_mb = 512
        ctx.invoked_function_arn = (
            "arn:aws:lambda:us-east-1:123456789012:function:miatun"
        )
        ctx.aws_request_id = "test-aws-request-id"
        return ctx

    def test_ping_no_headers_required(self):
        response = index.handler(self.make_event("/miatun/ping"), self.make_context())
        self.assertEqual(response["statusCode"], 200)

    def test_export_rejects_missing_api_key(self):
        event = self.make_event("/miatun/export", method="POST")
        response = index.handler(event, self.make_context())
        self.assertEqual(response["statusCode"], 401)

    def test_export_rejects_wrong_api_key(self):
        event = self.make_event(
            "/miatun/export", method="POST", headers={"internal-api-key": "wrong"}
        )
        response = index.handler(event, self.make_context())
        self.assertEqual(response["statusCode"], 401)

    def test_export_success(self):
        mock_dynamo = MagicMock()
        mock_dynamo.get_all.side_effect = [
            # received items scan
            [
                {
                    "id": {"S": "i1"},
                    "item_name": {"S": "Stand Mixer"},
                    "brand": {"S": "KitchenAid"},
                    "price_cents": {"N": "44999"},
                    "received": {"BOOL": True},
                    "claimant_id": {"S": "john_smith"},
                }
            ],
            # claims scan
            [
                {
                    "id": {"S": "c1"},
                    "item_id": {"S": "i1"},
                    "claimant_id": {"S": "john_smith"},
                    "claim_state": {"S": "PURCHASED"},
                }
            ],
        ]
        mock_dynamo.batch_get_items.return_value = [
            {
                "first_last": {"S": "john_smith"},
                "guest_pair_first_last": {"S": "jane_smith"},
                "phone": {"S": "+15555550100"},
                "street": {"S": "123 Main St"},
                "second_line": {"S": ""},
                "city": {"S": "Brooklyn"},
                "state_loc": {"S": "NY"},
                "zipcode": {"S": "11201"},
                "country": {"S": "USA"},
            }
        ]
        mock_s3 = MagicMock()
        mock_s3.generate_presigned_url.return_value = "https://example.com/signed"

        event = self.make_event(
            "/miatun/export",
            method="POST",
            headers={"Internal-Api-Key": "test_internal_key"},
        )

        with patch.object(index, "DYNAMO_CLIENT", mock_dynamo), patch.object(
            index, "S3_CLIENT", mock_s3
        ):
            response = index.handler(event, self.make_context())

        self.assertEqual(response["statusCode"], 200)
        body = json.loads(response["body"])
        self.assertEqual(body["message"], "export success")
        self.assertEqual(body["download_url"], "https://example.com/signed")
        self.assertEqual(body["row_count"], 1)

        # The uploaded object is a real xlsx with our data in it
        put_kwargs = mock_s3.put_object.call_args.kwargs
        self.assertEqual(put_kwargs["Bucket"], "test_export_bucket")
        self.assertTrue(put_kwargs["Key"].startswith("thank-you-cards/"))
        self.assertTrue(put_kwargs["Key"].endswith(".xlsx"))
        sheet = load_workbook(io.BytesIO(put_kwargs["Body"])).active
        self.assertEqual(sheet[2][0].value, "John Smith")
        self.assertEqual(sheet[2][2].value, "Stand Mixer")

    def test_export_returns_500_on_dynamo_error(self):
        mock_dynamo = MagicMock()
        mock_dynamo.get_all.side_effect = Exception("dynamo exploded")

        event = self.make_event(
            "/miatun/export",
            method="POST",
            headers={"Internal-Api-Key": "test_internal_key"},
        )

        with patch.object(index, "DYNAMO_CLIENT", mock_dynamo):
            response = index.handler(event, self.make_context())

        self.assertEqual(response["statusCode"], 500)

    def test_unknown_route_404(self):
        event = self.make_event(
            "/miatun/nope", headers={"x-first-last": "john_smith"}
        )
        response = index.handler(event, self.make_context())
        self.assertEqual(response["statusCode"], 404)


if __name__ == "__main__":
    unittest.main()
